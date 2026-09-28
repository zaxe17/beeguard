# services/backup_service.py
#
# Backup & Restore (admin → More → Backup & Restore).
#
# A backup is a plain .sql file in server/backups/ holding every table of
# the BeeGuard database (structure + rows). It's written with PyMySQL
# itself, so it works on Windows without mysqldump being installed.
#
#   beeguard_2026-09-28_13-20-05_manual.sql       — "Back Up Data Now"
#   beeguard_2026-09-29_02-00-00_auto.sql         — automatic (daily)
#   beeguard_2026-09-30_09-10-00_pre-restore.sql  — taken right BEFORE a
#                                                    restore, so a restore
#                                                    can always be undone
#
# Every backup ALSO gets an Excel copy with the same name (.xlsx): a
# "Summary" sheet + one sheet per table, for opening / sharing / checking
# the data. The download button gives the Excel file. Restore always uses
# the .sql file (Excel can't hold table structure, so it can't restore).
# Passwords / codes are hidden in the Excel copy (they stay in the .sql).
#
# Restore replaces ALL current data with the backup's. Uploaded images
# (report photos, chat photos, verification documents) are files, not
# database rows, so they are not part of the backup.

import datetime as dt
import json
import os
import re
import threading
import time
from decimal import Decimal

from pymysql.converters import escape_item

from config.database import Database

BACKEND_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_FOLDER = os.path.join(BACKEND_ROOT, "backups")

KINDS = ("manual", "auto", "pre-restore")
NAME_RE = re.compile(
    r"^beeguard_(\d{4}-\d{2}-\d{2})_(\d{2}-\d{2}-\d{2})_(manual|auto|pre-restore)\.sql$"
)
# Each statement in the file ends with this line, so restore can split
# statements safely even when saved text contains ";" or new lines.
END_MARK = "-- @@END@@"
INSERT_BATCH = 200
KEEP_AUTO = 7              # keep the newest 7 automatic backups
AUTO_EVERY_HOURS = 24
_lock = threading.Lock()   # one backup / restore at a time


def _sql_value(value) -> str:
    """
    A Python value as SQL text for the INSERT lines ('it\\'s', 1.50, NULL,
    '2026-09-01 08:00:00'...). Uses PyMySQL's own escaping directly —
    the pooled connection (DBUtils SteadyDBConnection) has no .escape().
    """
    return escape_item(value, "utf8mb4")


# ── Excel copy ────────────────────────────────
EXCEL_MAX_CELL = 32767                  # Excel's limit per cell
SECRET_COLUMNS = {"password", "code_hash", "p256dh", "auth", "endpoint"}


def _excel_name(sql_name: str) -> str:
    return sql_name[:-4] + ".xlsx"


def _excel_value(value, column: str):
    """A database value as something Excel can show."""
    if value is None:
        return None
    if column.lower() in SECRET_COLUMNS:
        return "(hidden)"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (bytes, bytearray)):
        return f"(binary, {len(value)} bytes)"
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_MAX_CELL]
    return value


def _write_excel(path: str, db_name: str, kind: str, tables: list[tuple]) -> None:
    """
    tables: [(table_name, columns, rows)]. Summary sheet first, then one
    sheet per table (bold yellow header, frozen, filter, sized columns).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="4A2F00")
    header_fill = PatternFill("solid", fgColor="FFDB4F")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["BeeGuard backup"])
    summary["A1"].font = Font(bold=True, size=14, color="4A2F00")
    summary.append(["Database", db_name])
    summary.append(["Type", kind])
    summary.append(["Created", dt.datetime.now().strftime("%B %d, %Y %I:%M %p")])
    summary.append([])
    summary.append(["Table", "Rows"])
    for cell in summary[6]:
        cell.font, cell.fill = header_font, header_fill
    for name, _cols, rows in tables:
        summary.append([name, len(rows)])
    summary.append(["Total", sum(len(r) for _n, _c, r in tables)])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 26

    used = {"Summary"}
    for name, columns, rows in tables:
        title = name[:31]
        n = 2
        while title in used:  # Excel sheet names: max 31 chars, unique
            title = f"{name[:28]}~{n}"
            n += 1
        used.add(title)

        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([_excel_value(row[c], c) for c in columns])

        # Date/time columns shown as dates, not numbers.
        for idx, col in enumerate(columns, start=1):
            sample = next((r[col] for r in rows if r[col] is not None), None)
            letter = get_column_letter(idx)
            if isinstance(sample, dt.datetime):
                for cell in ws[letter][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(sample, dt.date):
                for cell in ws[letter][1:]:
                    cell.number_format = "yyyy-mm-dd"
            # Width from the header + first 100 values (10–50 chars).
            longest = max(
                [len(str(col))] + [len(str(ws.cell(row=i, column=idx).value or ""))
                                   for i in range(2, min(ws.max_row, 101) + 1)]
            )
            ws.column_dimensions[letter].width = max(10, min(50, longest + 2))

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


# ── Excel from an existing .sql backup ────────
# Backups made before the Excel copy existed (or when openpyxl wasn't
# installed yet) only have the .sql. The first time one is downloaded,
# its Excel copy is built from the .sql file itself (our own format:
# one statement per END_MARK, values escaped by PyMySQL).
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}(\.\d+)?$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_UNESCAPE = {"0": "\0", "n": "\n", "r": "\r", "Z": "\x1a", "t": "\t", "b": "\b"}


def _parse_values(text: str) -> list[list]:
    """(1, 'it\\'s', NULL, 1.50), (...) -> [[1, "it's", None, 1.5], ...]"""
    rows, row, i, n = [], None, 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "(" and row is None:
            row = []
            i += 1
        elif ch == ")" and row is not None:
            rows.append(row)
            row = None
            i += 1
        elif ch in ", \n\r\t":
            i += 1
        elif row is None:
            i += 1
        elif ch == "'" or text.startswith("_binary'", i):
            binary = ch != "'"
            i += 8 if binary else 1
            out = []
            while i < n:
                c = text[i]
                if c == "\\" and i + 1 < n:
                    nxt = text[i + 1]
                    out.append(_UNESCAPE.get(nxt, nxt))
                    i += 2
                elif c == "'":
                    if i + 1 < n and text[i + 1] == "'":  # '' = one quote
                        out.append("'")
                        i += 2
                    else:
                        i += 1
                        break
                else:
                    out.append(c)
                    i += 1
            value = "".join(out)
            if binary:
                value = f"(binary, {len(value)} bytes)"
            elif _DATETIME_RE.match(value):
                try:
                    value = dt.datetime.fromisoformat(value)
                except ValueError:
                    pass
            elif _DATE_RE.match(value):
                try:
                    value = dt.date.fromisoformat(value)
                except ValueError:
                    pass
            row.append(value)
        else:  # NULL / number
            j = i
            while j < n and text[j] not in ",)":
                j += 1
            token = text[i:j].strip()
            i = j
            if token.upper() == "NULL":
                row.append(None)
            else:
                try:
                    row.append(int(token))
                except ValueError:
                    try:
                        row.append(float(token))
                    except ValueError:
                        row.append(token)
    return rows


def _tables_from_sql(text: str) -> tuple[str, list[tuple]]:
    """Reads our .sql backup back into [(table, columns, rows as dicts)]."""
    db_match = re.search(r"of `([^`]+)`", text)
    db_name = db_match.group(1) if db_match else "beeguard_system"
    order, columns, rows = [], {}, {}
    create_re = re.compile(r"CREATE TABLE `([^`]+)`")
    insert_re = re.compile(r"INSERT INTO `([^`]+)` \((.*?)\) VALUES\n(.*)\Z", re.S)
    for chunk in text.split(END_MARK):
        stmt = "\n".join(l for l in chunk.strip().splitlines()
                         if not l.strip().startswith("--")).strip()
        m = create_re.match(stmt)
        if m:
            table = m.group(1)
            order.append(table)
            columns[table] = re.findall(r"^\s+`([^`]+)`", stmt, re.M)
            rows[table] = []
            continue
        m = insert_re.match(stmt)
        if m:
            table = m.group(1)
            cols = [c.strip().strip("`") for c in m.group(2).split(",")]
            columns[table] = cols
            rows.setdefault(table, []).extend(
                dict(zip(cols, values)) for values in _parse_values(m.group(3))
            )
    return db_name, [(t, columns.get(t, []), rows.get(t, [])) for t in order]


def _now_name(kind: str) -> str:
    return f"beeguard_{dt.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}_{kind}.sql"


def _info(name: str) -> dict | None:
    m = NAME_RE.match(name)
    path = os.path.join(BACKUP_FOLDER, name)
    if not m or not os.path.isfile(path):
        return None
    created = dt.datetime.strptime(f"{m.group(1)} {m.group(2)}", "%Y-%m-%d %H-%M-%S")
    excel_path = os.path.join(BACKUP_FOLDER, _excel_name(name))
    has_excel = os.path.isfile(excel_path)
    return {
        "name": name,
        "kind": m.group(3),
        "created_at": created.isoformat(),
        "size_bytes": os.path.getsize(path),
        "has_excel": has_excel,
        "excel_name": _excel_name(name) if has_excel else None,
        "excel_size_bytes": os.path.getsize(excel_path) if has_excel else None,
    }


def _path_for(name: str) -> str:
    """Full path of an existing backup. Only our own file names are allowed."""
    if not NAME_RE.match(name or ""):
        raise LookupError("Backup not found.")
    path = os.path.join(BACKUP_FOLDER, name)
    if not os.path.isfile(path):
        raise LookupError("Backup not found.")
    return path


class BackupService:

    # ── LIST ──────────────────────────────────
    @staticmethod
    def list_backups() -> list[dict]:
        if not os.path.isdir(BACKUP_FOLDER):
            return []
        items = [i for i in (_info(n) for n in os.listdir(BACKUP_FOLDER)) if i]
        items.sort(key=lambda i: i["created_at"], reverse=True)
        return items

    @staticmethod
    def latest() -> dict | None:
        items = BackupService.list_backups()
        return items[0] if items else None

    @staticmethod
    def file_path(name: str, fmt: str = "sql") -> tuple[str, str]:
        """
        (path, download name). fmt "xlsx" = the Excel copy — made from the
        .sql right now if this backup doesn't have one yet.
        """
        path = _path_for(name)
        if fmt == "xlsx":
            excel = os.path.join(BACKUP_FOLDER, _excel_name(name))
            if not os.path.isfile(excel):
                with _lock:
                    if not os.path.isfile(excel):
                        with open(path, encoding="utf-8") as f:
                            db_name, tables = _tables_from_sql(f.read())
                        kind = NAME_RE.match(name).group(3)
                        _write_excel(excel, db_name, kind, tables)
            return excel, _excel_name(name)
        return path, name

    # ── CREATE ────────────────────────────────
    @staticmethod
    def create_backup(kind: str = "manual") -> dict:
        if kind not in KINDS:
            raise ValueError("Invalid backup kind.")
        with _lock:
            return BackupService._write_backup(kind)

    @staticmethod
    def _write_backup(kind: str) -> dict:
        os.makedirs(BACKUP_FOLDER, exist_ok=True)
        name = _now_name(kind)
        path = os.path.join(BACKUP_FOLDER, name)
        tmp = path + ".tmp"

        conn = Database.get_connection()
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT DATABASE() AS db")
                db_name = (cur.fetchone() or {}).get("db")
                cur.execute("SHOW FULL TABLES WHERE Table_type = 'BASE TABLE'")
                tables = [list(r.values())[0] for r in cur.fetchall()]

            excel_tables = []  # (table, columns, rows) for the Excel copy
            with open(tmp, "w", encoding="utf-8") as f:
                f.write(f"-- BeeGuard backup ({kind}) of `{db_name}`\n")
                f.write(f"-- Created {dt.datetime.now().isoformat(timespec='seconds')}\n")
                f.write(f"-- Tables: {len(tables)}\n\n")
                f.write(f"SET FOREIGN_KEY_CHECKS = 0\n{END_MARK}\n")

                for table in tables:
                    with conn.cursor() as cur:
                        cur.execute(f"SHOW CREATE TABLE `{table}`")
                        create_sql = list(cur.fetchone().values())[1]
                    f.write(f"\n-- Table `{table}`\n")
                    f.write(f"DROP TABLE IF EXISTS `{table}`\n{END_MARK}\n")
                    f.write(f"{create_sql}\n{END_MARK}\n")

                    with conn.cursor() as cur:
                        cur.execute(f"SELECT * FROM `{table}`")
                        columns = [d[0] for d in cur.description]
                        col_list = ", ".join(f"`{c}`" for c in columns)
                        batch = []
                        rows = cur.fetchall()
                        excel_tables.append((table, columns, rows))
                        for row in rows:
                            values = ", ".join(_sql_value(row[c]) for c in columns)
                            batch.append(f"({values})")
                            if len(batch) >= INSERT_BATCH:
                                f.write(f"INSERT INTO `{table}` ({col_list}) VALUES\n"
                                        + ",\n".join(batch) + f"\n{END_MARK}\n")
                                batch = []
                        if batch:
                            f.write(f"INSERT INTO `{table}` ({col_list}) VALUES\n"
                                    + ",\n".join(batch) + f"\n{END_MARK}\n")

                f.write(f"\nSET FOREIGN_KEY_CHECKS = 1\n{END_MARK}\n")
            os.replace(tmp, path)
        except Exception:
            if os.path.exists(tmp):
                os.remove(tmp)
            raise
        finally:
            conn.close()

        # Excel copy — best-effort: the .sql backup is what counts.
        try:
            _write_excel(os.path.join(BACKUP_FOLDER, _excel_name(name)),
                         db_name, kind, excel_tables)
        except Exception as e:
            print(f"[BACKUP] Excel copy failed for {name}: {e}")

        if kind == "auto":
            BackupService._prune_auto()
        return _info(name)

    @staticmethod
    def _prune_auto() -> None:
        autos = [b for b in BackupService.list_backups() if b["kind"] == "auto"]
        for old in autos[KEEP_AUTO:]:
            for file_name in (old["name"], _excel_name(old["name"])):
                try:
                    os.remove(os.path.join(BACKUP_FOLDER, file_name))
                except OSError:
                    pass

    # ── RESTORE ───────────────────────────────
    @staticmethod
    def restore(name: str) -> dict:
        """
        Replaces all current data with the backup's. A 'pre-restore'
        backup of the current data is taken first, so this can be undone
        by restoring that one.
        """
        path = _path_for(name)
        with _lock:
            safety = BackupService._write_backup("pre-restore")

            with open(path, encoding="utf-8") as f:
                text = f.read()
            statements = [
                s.strip() for s in text.split(END_MARK)
                if s.strip() and not all(
                    line.strip().startswith("--") or not line.strip()
                    for line in s.strip().splitlines()
                )
            ]

            conn = Database.get_connection()
            try:
                with conn.cursor() as cur:
                    for stmt in statements:
                        cur.execute(stmt)
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                try:
                    with conn.cursor() as cur:
                        cur.execute("SET FOREIGN_KEY_CHECKS = 1")
                except Exception:
                    pass
                conn.close()

        return {"restored": name, "safety_backup": safety}

    # ── AUTOMATIC (daily) ─────────────────────
    @staticmethod
    def run_auto_backup_if_due() -> dict | None:
        """Makes an 'auto' backup when automatic backup is on and none was made in 24h."""
        from services.settings_service import SettingsService  # avoid import cycle
        if not SettingsService.get_system().get("auto_backup"):
            return None
        last_auto = next(
            (b for b in BackupService.list_backups() if b["kind"] == "auto"), None
        )
        if last_auto:
            age = dt.datetime.now() - dt.datetime.fromisoformat(last_auto["created_at"])
            if age < dt.timedelta(hours=AUTO_EVERY_HOURS):
                return None
        return BackupService.create_backup("auto")


_auto_check_lock = threading.Lock()
_next_auto_check = 0.0
AUTO_CHECK_EVERY_S = 30 * 60


def maybe_run_auto_backup() -> None:
    """
    Called on every request (app.py before_request) but only does work
    once every 30 minutes: if automatic backup is on and the last one is
    older than 24 hours, a new one is made in a background thread (the
    request isn't slowed down).

    Checked on requests instead of a timer thread because in debug mode
    Flask runs two processes (reloader + app) and a timer would back up
    twice — only the real app process receives requests.
    """
    global _next_auto_check
    now = time.monotonic()
    if now < _next_auto_check:
        return
    with _auto_check_lock:
        if now < _next_auto_check:
            return
        _next_auto_check = now + AUTO_CHECK_EVERY_S

    def run():
        try:
            made = BackupService.run_auto_backup_if_due()
            if made:
                print(f"[BACKUP] Automatic backup saved: {made['name']}")
        except Exception as e:
            print(f"[BACKUP] Automatic backup failed: {e}")

    threading.Thread(target=run, name="beeguard-auto-backup", daemon=True).start()